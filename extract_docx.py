import glob
import docx


def dump(pattern, out):
    matches = glob.glob(pattern)
    if not matches:
        print("NO MATCH:", pattern)
        return
    src = matches[0]
    d = docx.Document(src)
    text = "\n".join(p.text for p in d.paragraphs)
    with open(out, "w", encoding="utf-8") as f:
        f.write(text)
    print(f"{out} <- {src} ({len(d.paragraphs)} paras, {len(text)} chars)")


dump(r"C:\Users\asus\Desktop\*Linux*.docx", r"C:\Users\asus\Desktop\git\t8.txt")
